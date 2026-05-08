"""Excel Parser"""
from typing import List, Dict, Tuple

class ExcelParser:
    @staticmethod
    def parse_xlsx(file_path: str) -> Tuple[List[str], List[Dict]]:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path)
            ws = wb.active
            
            headers = [cell.value for cell in ws[1]]
            rows = []
            
            for row in ws.iter_rows(min_row=2, values_only=False):
                row_dict = {}
                for i, cell in enumerate(row):
                    if i < len(headers):
                        row_dict[headers[i]] = cell.value
                rows.append(row_dict)
            
            return headers, rows
        except ImportError:
            return [], []
    
    @staticmethod
    def parse_xls(file_path: str) -> Tuple[List[str], List[Dict]]:
        try:
            import xlrd
            wb = xlrd.open_workbook(file_path)
            ws = wb.sheet_by_index(0)
            
            headers = [ws.cell_value(0, i) for i in range(ws.ncols)]
            rows = []
            
            for row_idx in range(1, ws.nrows):
                row_dict = {}
                for col_idx in range(ws.ncols):
                    row_dict[headers[col_idx]] = ws.cell_value(row_idx, col_idx)
                rows.append(row_dict)
            
            return headers, rows
        except ImportError:
            return [], []